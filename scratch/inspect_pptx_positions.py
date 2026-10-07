import pptx

prs = pptx.Presentation(r'C:\Users\praja\Downloads\Project Poster-Template-31-07-2026.pptx')
slide = prs.slides[0]

boxes = []
for idx, shape in enumerate(slide.shapes):
    if shape.has_text_frame:
        boxes.append({
            'idx': idx,
            'id': shape.shape_id,
            'name': shape.name,
            'left': round(shape.left.inches, 2),
            'top': round(shape.top.inches, 2),
            'width': round(shape.width.inches, 2),
            'height': round(shape.height.inches, 2),
            'text': shape.text_frame.text.strip().replace('\n', ' ')[:60]
        })

# Sort by Left, then Top
boxes.sort(key=lambda x: (x['left'], x['top']))

for b in boxes:
    print(f"L={b['left']:5.2f}\", T={b['top']:5.2f}\" | W={b['width']:5.2f}\", H={b['height']:5.2f}\" | ID={b['id']:2d} ({b['name']:15s}) -> \"{b['text']}\"")
