import pptx

prs = pptx.Presentation(r'C:\Users\praja\Downloads\Project Poster-Template-31-07-2026.pptx')
slide = prs.slides[0]

def inspect_shape(shape, depth=0):
    indent = "  " * depth
    print(f"{indent}ID: {shape.shape_id} | Name: '{shape.name}' | Type: {shape.shape_type}")
    if shape.has_text_frame:
        txt = shape.text_frame.text.strip().replace('\n', ' -- ')
        print(f"{indent}   Text: \"{txt[:120]}\"")
    if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.GROUP:
        print(f"{indent}   [GROUP CONTAINS {len(shape.shapes)} SUB-SHAPES]:")
        for sub in shape.shapes:
            inspect_shape(sub, depth+1)

for idx, shape in enumerate(slide.shapes):
    print(f"\nShape #{idx}:")
    inspect_shape(shape)
