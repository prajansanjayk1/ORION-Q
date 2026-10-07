import pptx

prs = pptx.Presentation(r'C:\Users\praja\Downloads\Project Poster-Template-31-07-2026.pptx')
slide = prs.slides[0]

print(f"Slide Dimensions: {prs.slide_width.inches:.2f} x {prs.slide_height.inches:.2f} inches")
print("="*60)

for idx, shape in enumerate(slide.shapes):
    print(f"\n--- Shape #{idx} | ID: {shape.shape_id} | Name: '{shape.name}' | Type: {shape.shape_type} ---")
    print(f"    Bounding Box: Left={shape.left.inches:.2f}\", Top={shape.top.inches:.2f}\", Width={shape.width.inches:.2f}\", Height={shape.height.inches:.2f}\"")
    if shape.has_text_frame:
        print("    Text Content:")
        for p_idx, p in enumerate(shape.text_frame.paragraphs):
            print(f"      P{p_idx}: \"{p.text}\"")
    if shape.has_table:
        print(f"    Table ({len(shape.table.rows)} rows x {len(shape.table.columns)} cols):")
        for r_i, row in enumerate(shape.table.rows):
            row_txt = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
            print(f"      Row {r_i}: {row_txt}")
