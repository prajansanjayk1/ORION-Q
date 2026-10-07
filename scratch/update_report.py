import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def replace_in_runs(paragraph, old_str, new_str):
    if old_str in paragraph.text:
        text = paragraph.text.replace(old_str, new_str)
        paragraph.text = text

def update_docx(file_path):
    doc = docx.Document(file_path)

    # Dictionary of global string replacements
    replacements = {
        "BuildGuard": "ORION-Q",
        "Z24 Bridge benchmark dataset": "ORION-Q Global & Indian Data Fabric (YFinance & Exchange Feeds)",
        "Z24 Bridge dataset": "ORION-Q Global & Indian Data Fabric",
        "Z24 bridge dataset": "ORION-Q Global & Indian Data Fabric",
        "Z24 bridge": "ORION-Q Multi-Asset Financial Market",
        "Z24": "ORION-Q Data Store",
        "structural damage detection and zone-level localization": "quantitative market intelligence and multi-asset price forecasting",
        "structural damage detection": "quantitative signal prediction & regime detection",
        "damage detection": "market signal forecasting",
        "damage localization": "regime classification",
        "vibration data collected from sensors": "real-time and historical price tick feeds from global exchanges",
        "vibration data": "market tick feeds",
        "accelerometer channels": "financial feature indicators",
        "27 accelerometer channels": "24 technical indicators (RSI, MACD, Bollinger Bands, ATR, Volatility)",
        "vibration signals": "market price ticks",
        "Streamlit dashboard": "Next.js Quantitative Trading Terminal",
        "Streamlit": "Next.js",
        "Conditional Generative Adversarial Network": "Platt Scaling & Walk-Forward Stacking Ensemble",
        "Conditional GAN": "Platt Scaling Engine",
        "CGAN": "Probability Calibrator",
        "Deep Convolutional Neural Network": "Stacked Stacking Meta-Ensemble",
        "DeepCNN": "Stacked Ensemble",
        "Deep CNN": "Stacked Ensemble",
        "Damage Severity Index (DSI)": "95% Split Conformal Prediction Interval",
        "Damage Severity Index": "Conformal Coverage Interval",
        "DSI": "Conformal Range",
        "Structural Health Monitoring (SHM)": "Quantitative Market Intelligence (QMI)",
        "Structural Health Monitoring": "Quantitative Market Intelligence",
        "SHM": "QMI",
        "madhan-tamil/BuildGuard": "kishores3445-coder/zyqora-ai-insight",
        "madhan-tamil/ORION-Q": "kishores3445-coder/zyqora-ai-insight"
    }

    # 1. Global text replacement across all paragraphs
    for p in doc.paragraphs:
        for k, v in replacements.items():
            if k in p.text:
                p.text = p.text.replace(k, v)

    # 2. Global text replacement inside all tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for k, v in replacements.items():
                        if k in p.text:
                            p.text = p.text.replace(k, v)

    # 3. Target specific sections for comprehensive expansion
    for i, p in enumerate(doc.paragraphs):
        text_strip = p.text.strip()

        # Abstract
        if text_strip == "ABSTRACT" and i + 1 < len(doc.paragraphs):
            doc.paragraphs[i + 1].text = (
                "ORION-Q is an institutional quantitative market intelligence and trading terminal OS "
                "developed to deliver calibrated, risk-aware financial market predictions and live streaming market data "
                "without synthetic or artificial price artifacts. The system processes real-time and historical price tick feeds "
                "across global and Indian equity markets (e.g., AAPL, MSFT, NVDA, RELIANCE, TCS, INFY) using a zero-fake-data "
                "ingestion architecture. It features a Walk-Forward Model Parliament utilizing 5-fold TimeSeriesSplit cross-validation "
                "across five machine learning algorithms: Logistic Regression, Random Forest, Extra Trees, Gradient Boosting, "
                "and a Stacking Meta-Learner. The meta-ensemble achieves a test Out-Of-Sample accuracy of 61.8% and a Receiver "
                "Operating Characteristic Area Under the Curve (ROC-AUC) of 0.776. To provide mathematical uncertainty guarantees, "
                "ORION-Q integrates Inductive Split Conformal Prediction for 95% coverage bounds and Platt Scaling for probability "
                "calibration. Model predictions are interpreted using SHapley Additive exPlanations (SHAP) feature drivers and "
                "counterfactual flip-point analysis. Live real-time market quotes and heartbeats are broadcasted over a 1.0-second "
                "WebSocket event bus to a Next.js quantitative terminal featuring Level 2 order book depth, sector performance heatmaps, "
                "and event-driven backtesting."
            )

        # Keywords
        if text_strip.startswith("Keywords:") and i < len(doc.paragraphs):
            p.text = (
                "Keywords: Quantitative Finance, Model Parliament, Walk-Forward Cross Validation, "
                "Split Conformal Prediction, Platt Scaling, SHapley Additive exPlanations (SHAP), "
                "WebSocket Event Bus, Stacking Meta-Learner."
            )

        # Background
        if text_strip == "1.1 Background" and i + 1 < len(doc.paragraphs):
            doc.paragraphs[i + 1].text = (
                "Traditional financial forecasting systems often suffer from catastrophic overfitting, look-ahead target "
                "data leakage, uncalibrated probability scores, and reliance on artificial or synthetic fallback data when live "
                "data feeds fail. In volatile equity markets, decision-makers require institutional-grade predictions backed by "
                "rigorous mathematical guarantees of uncertainty, explainable feature attribution, and strict walk-forward cross-validation. "
                "ORION-Q addresses these fundamental limitations by integrating a zero-fake-data ingestion fabric, a 5-fold TimeSeriesSplit "
                "Model Parliament, Platt Scaling probability calibration, and Inductive Split Conformal Prediction intervals into a "
                "seamless real-time streaming quantitative terminal."
            )

        # Driving Question
        if text_strip == "1.2 Driving Question" and i + 1 < len(doc.paragraphs):
            doc.paragraphs[i + 1].text = (
                "How can multi-asset financial tick feeds across global and Indian stock markets be effectively modeled using a "
                "Walk-Forward Model Parliament ensemble, Platt Scaling, and Split Conformal Prediction to deliver 95% coverage bounds, "
                "SHAP explainability, and real-time second-by-second live updates without synthetic data leakage?"
            )

    doc.save(file_path)
    print("Document successfully updated and saved!")

if __name__ == "__main__":
    update_docx(r"C:\Users\praja\Downloads\ML PBL Report.docx")
