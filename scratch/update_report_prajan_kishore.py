import docx
import os
import shutil

def update_prajan_kishore_report(doc_path, output_path):
    doc = docx.Document(doc_path)
    print(f"Loaded document with {len(doc.paragraphs)} paragraphs and {len(doc.tables)} tables.")

    # 1. Update Title Page Title
    doc.paragraphs[0].text = (
        "ORION-Q: INSTITUTIONAL QUANTITATIVE MARKET OS WITH WALK-FORWARD MODEL PARLIAMENT, "
        "SPLIT CONFORMAL PREDICTION, PLATT-STYLE PROBABILITY CALIBRATION, SHAP EXPLAINABILITY, AND LIVE MARKET EVENT STREAMING"
    )

    # 2. Update Student Names Across Paragraphs
    name_replacements = {
        "MADHAN T": "PRAJAN SANJAY K",
        "AVINASH N T": "KISHORE S",
        "Madhan T": "Prajan Sanjay K",
        "Avinash N T": "Kishore S",
        "Madhan": "Prajan Sanjay K",
        "Avinash": "Kishore S",
        "210425149028": "",
        "210425149006": "",
        "Platt Scaling Calibration": "Platt-Style Probability Calibration"
    }

    # 3. Update Bonafide Certificate (P[042])
    for p in doc.paragraphs:
        if "BONAFIDE CERTIFICATE" in p.text or "Project–Based Learning report titled" in p.text or "Bonafide record of work carried out by" in p.text:
            p.text = (
                "This is to certify that the Project–Based Learning report titled “ORION-Q: Institutional Quantitative Market OS "
                "with Walk-Forward Model Parliament, Split Conformal Prediction, Platt-Style Probability Calibration, SHAP Explainability, "
                "and Live Market Event Streaming” is a Bonafide record of work carried out by Prajan Sanjay K, Kishore S of the "
                "Department of Computer Science and Engineering (Cyber Security), Chennai Institute of Technology, as part of the "
                "continuous, mentor–guided Project-Based Learning (PBL) component of the Machine Learning course during the academic year [2026–2027] under my supervision."
            )

    # 4. Update Declaration (P[055])
    for p in doc.paragraphs:
        if "We jointly declare that the PBL report on" in p.text:
            p.text = (
                "We jointly declare that the PBL report on “ORION-Q: Institutional Quantitative Market OS with Walk-Forward Model Parliament, "
                "Split Conformal Prediction, Platt-Style Probability Calibration, SHAP Explainability, and Live Market Event Streaming” "
                "is the result of original work done by us and best of our knowledge, similar work has not been submitted to ANNA UNIVERSITY, CHENNAI "
                "for the requirement of Degree of BACHELOR OF ENGINEERING. This PBL report is submitted on the partial fulfilment of the requirement "
                "of the award of Degree of COMPUTER SCIENCE AND ENGINEERING (CYBER SECURITY)."
            )

    # 5. Global Student Name Replacements in all Paragraphs
    for p in doc.paragraphs:
        for old_n, new_n in name_replacements.items():
            if old_n in p.text:
                p.text = p.text.replace(old_n, new_n)

    # 6. Global Student Name Replacements in all Tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for old_n, new_n in name_replacements.items():
                        if old_n in p.text:
                            p.text = p.text.replace(old_n, new_n)

    # 7. Clean up extra hyphens or trailing IDs in student names
    for p in doc.paragraphs:
        if "PRAJAN SANJAY K –," in p.text:
            p.text = p.text.replace("PRAJAN SANJAY K –,", "PRAJAN SANJAY K,")
        if "Prajan Sanjay K –," in p.text:
            p.text = p.text.replace("Prajan Sanjay K –,", "Prajan Sanjay K,")
        if "Prajan Sanjay K –" in p.text:
            p.text = p.text.replace("Prajan Sanjay K –", "Prajan Sanjay K")

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if "Prajan Sanjay K –" in p.text:
                        p.text = p.text.replace("Prajan Sanjay K –", "Prajan Sanjay K")

    doc.save(output_path)
    print(f"Report successfully updated for Prajan Sanjay K & Kishore S and saved to {output_path}!")

if __name__ == "__main__":
    src_file = r"C:\Users\praja\Downloads\ORION-Q_PBL_Report_Final_With_Figures.docx"
    output_file = r"C:\Users\praja\Downloads\ORION-Q_PBL_Report_Prajan_Kishore.docx"
    update_prajan_kishore_report(src_file, output_file)
